"""Tests for 0_tools/bugs.py -- always against a temporary copy of
bugsDB.xlsx, never the real one.

Run with:  pytest 0_tests/ -v
"""

import datetime
import os
import shutil
import sys

import openpyxl
import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJECT_ROOT, '0_tools'))
import bugs  # noqa: E402


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / 'bugsDB.xlsx'
    shutil.copy(os.path.join(PROJECT_ROOT, 'bugsDB.xlsx'), path)
    return str(path)


def run(workbook, *argv):
    bugs.main(['--file', workbook, *argv])


def lastBug(workbook):
    log = bugs.BugLog(workbook)
    row = log.dataRows()[-1]
    return log, row


def test_drop_down_choices_are_read_from_the_sheet(workbook):
    log = bugs.BugLog(workbook)
    assert log.choices['Status'] == ['Open', 'Fixed']
    assert log.choices['Severity'] == ['Low', 'Medium', 'High']
    assert log.choices['Server'] == ['DEV', 'LIVE']


def test_add_assigns_next_id_defaults_and_template_style(workbook):
    before = bugs.BugLog(workbook)
    expectedID = before.nextID()
    templateRow = before.dataRows()[-1]

    run(workbook, 'add', '--title', 'Test bug', '--severity', 'medium', '--actual', 'Broken')

    log, row = lastBug(workbook)
    assert log.get(row, 'ID') == expectedID
    assert log.get(row, 'Title') == 'Test bug'
    assert log.get(row, 'Severity') == 'Medium'  # normalised to the drop-down's spelling
    assert log.get(row, 'Status') == 'Open'
    assert log.get(row, 'Bug / Feature Request') == 'Bug'
    assert log.get(row, 'Server') == 'DEV'
    assert log.get(row, 'Date Found') is not None
    dateColumn = log.columns['Date Found']
    assert log.sheet.cell(row, dateColumn).number_format == log.sheet.cell(templateRow, dateColumn).number_format
    assert os.path.exists(workbook + '.bak')


def test_add_live_server(workbook):
    run(workbook, 'add', '--title', 'Seen on live', '--server', 'live')
    log, row = lastBug(workbook)
    assert log.get(row, 'Server') == 'LIVE'


def test_add_fixed_sets_date_fixed(workbook):
    run(workbook, 'add', '--title', 'Already fixed', '--status', 'Fixed')
    log, row = lastBug(workbook)
    assert log.get(row, 'Status') == 'Fixed'
    assert log.get(row, 'Date Fixed') is not None


def test_add_rejects_a_value_not_in_the_drop_down(workbook):
    with pytest.raises(SystemExit):
        run(workbook, 'add', '--title', 'Bad', '--severity', 'Catastrophic')


def test_add_dry_run_does_not_save(workbook):
    count = len(bugs.BugLog(workbook).dataRows())
    run(workbook, 'add', '--title', 'Not saved', '--dry-run')
    assert len(bugs.BugLog(workbook).dataRows()) == count


def test_add_prompts_when_no_title_given(workbook, monkeypatch):
    answers = iter(['Prompted bug', '', 'High'] + [''] * 20)
    monkeypatch.setattr('builtins.input', lambda prompt='': next(answers))
    run(workbook, 'add')
    log, row = lastBug(workbook)
    assert log.get(row, 'Title') == 'Prompted bug'
    assert log.get(row, 'Status') == 'Open'
    assert log.get(row, 'Severity') == 'High'


def test_set_fixed_and_reopen(workbook):
    run(workbook, 'add', '--title', 'To fix')
    log, row = lastBug(workbook)
    bugID = str(log.get(row, 'ID'))

    run(workbook, 'set', bugID, '--fixed')
    log, row = lastBug(workbook)
    assert log.get(row, 'Status') == 'Fixed'
    assert log.get(row, 'Date Fixed') is not None

    run(workbook, 'set', bugID, '--reopen')
    log, row = lastBug(workbook)
    assert log.get(row, 'Status') == 'Open'
    assert log.get(row, 'Date Fixed') is None


def test_set_unknown_id_exits(workbook):
    with pytest.raises(SystemExit):
        run(workbook, 'set', '9999', '--fixed')


def test_validations_and_other_sheets_survive_a_save(workbook):
    before = openpyxl.load_workbook(workbook)
    validations = len(before['bugsDB'].data_validations.dataValidation)
    run(workbook, 'add', '--title', 'Check survival')
    after = openpyxl.load_workbook(workbook)
    assert after.sheetnames == before.sheetnames
    assert len(after['bugsDB'].data_validations.dataValidation) == validations


# --- import-reports ------------------------------------------------------

def fakeReport(server, reportID, title, dateFound='2026-10-17T15:09:38+00:00'):
    return {'id': reportID, 'dateFound': dateFound, 'title': title, 'bugOrFeature': 'Feature-Request',
            'severity': 'medium', 'component': 'GRD-TRACK', 'description': 'Something odd happened.',
            'stepsToReproduce': '1. Open GRD-TRACK', 'expectedBehaviour': None, 'actualBehaviour': None,
            'foundBy': '', 'server': server, 'phase': 'Alpha', 'userAgent': 'pytest',
            'viewport': '375x812', 'visitorHash': 'x', 'copiedToBugsDb': 0}


@pytest.fixture
def reports(monkeypatch):
    """Fake sources: {'LIVE': [...], 'DEV': [...]}; records what gets marked."""
    sources = {'LIVE': [], 'DEV': []}
    marked = {'LIVE': [], 'DEV': []}
    monkeypatch.setattr(bugs, 'fetchReports', lambda source: [r for r in sources[source] if r['id'] not in marked[source]])
    monkeypatch.setattr(bugs, 'markReportsCopied', lambda source, ids: marked[source].extend(ids))
    return sources, marked


def test_import_reports_copies_both_sources_oldest_first(workbook, reports):
    sources, marked = reports
    sources['LIVE'] = [fakeReport('LIVE', 4, 'Live one', '2026-10-17T16:00:00+00:00')]
    sources['DEV'] = [fakeReport('DEV', 1, 'Dev one', '2026-10-17T09:30:05+00:00')]
    firstID = bugs.BugLog(workbook).nextID()

    run(workbook, 'import-reports')

    log = bugs.BugLog(workbook)
    devRow, liveRow = log.dataRows()[-2:]
    assert [log.get(devRow, 'ID'), log.get(liveRow, 'ID')] == [firstID, firstID + 1]
    assert log.get(devRow, 'Title') == 'Dev one' and log.get(devRow, 'Server') == 'DEV'
    assert log.get(liveRow, 'Server') == 'LIVE'
    assert log.get(devRow, 'Date Found') == datetime.datetime(2026, 10, 17, 9, 30, 5)
    assert '[report DEV#1]' in log.get(devRow, 'Notes') and 'DTG 2026-10-17 09:30:05Z' in log.get(devRow, 'Notes')
    assert log.get(devRow, 'Severity') == 'Medium' and log.get(devRow, 'Phase') == 'Alpha'
    assert log.get(devRow, 'Status') == 'Open' and log.get(devRow, 'Found By') == 'Anonymous tester'
    assert marked == {'LIVE': [4], 'DEV': [1]}


def test_import_reports_never_imports_twice(workbook, reports, monkeypatch):
    sources, marked = reports
    sources['LIVE'] = [fakeReport('LIVE', 7, 'Once only')]
    monkeypatch.setattr(bugs, 'markReportsCopied', lambda source, ids: None)  # marking fails silently
    run(workbook, 'import-reports')
    count = len(bugs.BugLog(workbook).dataRows())
    run(workbook, 'import-reports')
    assert len(bugs.BugLog(workbook).dataRows()) == count


def test_import_reports_dry_run_saves_and_marks_nothing(workbook, reports):
    sources, marked = reports
    sources['DEV'] = [fakeReport('DEV', 2, 'Dry')]
    count = len(bugs.BugLog(workbook).dataRows())
    run(workbook, 'import-reports', '--dry-run')
    assert len(bugs.BugLog(workbook).dataRows()) == count
    assert marked == {'LIVE': [], 'DEV': []}
