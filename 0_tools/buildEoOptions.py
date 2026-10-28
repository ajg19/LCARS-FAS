"""Builds assets/eoOptions.js from Uhura's consolidated EO table: the option lists for the EO control
panel's ORGANISATION, SATELLITE and USE CASE drop-downs, and EO_SOURCES, the one-entry-per-source
table its SOURCE drop-down is filtered from.

    python 0_tools/buildEoOptions.py [path to the xlsx]

The workbook is a working document (enterpriseCrewOutput/ is never committed), so the generated
file is what is committed. Re-run this after the workbook's CSV columns change. LEVEL, FREQUENCY and
RESOLUTION are fixed vocabularies that live in assets/eoPanel.js, not here."""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, 'enterpriseCrewOutput', 'uhura_eoConsolidatedApiTable_2026-10-25.xlsx')
OUT = os.path.join(ROOT, 'assets', 'eoOptions.js')

# Header text of each CSV column in the workbook -> the constant it fills.
COLUMNS = {
    'Satellites (CSV)': 'EO_SATELLITES',
    'Organisations (CSV)': 'EO_ORGANISATIONS',
    'Use cases (CSV)': 'EO_USE_CASES',
}

# Catch-all phrases in the satellite column that are not names anyone could pick.
NOT_A_NAME = {'many others', 'other gnss-ro missions', 'other altimeter and ocean-colour missions'}


def splitCsv(cell):
    """Splits a CSV cell on commas that are not inside brackets (an earlier version of the
    workbook had 'Maxar constellation (WorldView, GeoEye)' in one cell)."""
    return [part.strip() for part in re.split(r',\s*(?![^()]*\))', str(cell)) if part.strip()]


def buildOptions(path):
    sheet = openpyxl.load_workbook(path, data_only=True)['EO API table']
    headers = {sheet.cell(1, c).value: c for c in range(1, sheet.max_column + 1)}
    options = {}
    for header, constant in COLUMNS.items():
        column = headers[header]
        values = set()
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column).value
            if cell:
                values.update(v for v in splitCsv(cell) if v.lower() not in NOT_A_NAME)
        options[constant] = sorted(values, key=str.lower)
    return options


SOURCE_FIELDS = {   # EO_SOURCES key -> workbook column header
    'levels': 'Levels available (CSV)',
    'bands': 'Methodologies (CSV)',
    'resolutions': 'Resolution ranges available, metres (CSV)',
    'organisations': 'Organisations (CSV)',
    'satellites': 'Satellites (CSV)',
    'useCases': 'Use cases (CSV)',
}
DATA_SOURCE_HEADER = 'Data source (organisation / archive)'
MODALITY_HEADER = 'Modality / Frequency'
LICENCE_OK_HEADER = 'Free API + licence OK (Y/N)'
LICENCE_WHY_HEADER = 'Why'
VISUALISATION_HEADER = 'Suggested visualisation (added 2026-10-28)'


def shortArchive(dataSource):
    """The archive's short name: the text after the last colon, before any semicolon."""
    return str(dataSource or '').split(';')[0].split(':')[-1].strip()


def sourceLabel(satellites, dataSource):
    """'Sentinel-2A, Sentinel-2B, Sentinel-2C - Copernicus Data Space Ecosystem (CDSE)': up to three
    satellites, then the archive's short name (the text after the last colon, before any semicolon)."""
    names = ', '.join(satellites[:3]) + (f' +{len(satellites) - 3}' if len(satellites) > 3 else '')
    archive = shortArchive(dataSource)
    return f'{names} - {archive}' if names and archive else names or archive


def buildSources(path):
    sheet = openpyxl.load_workbook(path, data_only=True)['EO API table']
    headers = {sheet.cell(1, c).value: c for c in range(1, sheet.max_column + 1)}
    sources = []
    for row in range(2, sheet.max_row + 1):
        source = {}
        for key, header in SOURCE_FIELDS.items():
            cell = sheet.cell(row, headers[header]).value
            values = splitCsv(cell) if cell else []
            source[key] = [v for v in values if v.lower() not in NOT_A_NAME]
        source['archive'] = shortArchive(sheet.cell(row, headers[DATA_SOURCE_HEADER]).value)
        source['label'] = sourceLabel(source['satellites'], sheet.cell(row, headers[DATA_SOURCE_HEADER]).value)
        source['modality'] = ' '.join(str(sheet.cell(row, headers[MODALITY_HEADER]).value or '').split())
        source['licenceOk'] = str(sheet.cell(row, headers[LICENCE_OK_HEADER]).value or '').strip() == 'Y'
        source['licence'] = ' '.join(str(sheet.cell(row, headers[LICENCE_WHY_HEADER]).value or '').split())
        # The class word before the colon (IMAGE OVERLAY, GRID MAP, TRACK PLOT, ...): only IMAGE OVERLAY sources can be drawn on the map so far.
        source['visualisation'] = str(sheet.cell(row, headers[VISUALISATION_HEADER]).value or '').split(':')[0].strip()
        sources.append(source)
    # Two sources can share satellites and archive (Sentinel-3's altimeter and its imagers): tell them apart by modality.
    counts = {}
    for source in sources:
        counts[source['label']] = counts.get(source['label'], 0) + 1
    for source in sources:
        if counts[source['label']] > 1 and source['modality']:
            source['label'] += f" ({source['modality'].split('(')[0].strip()})"
    return sources


def render(options, sources):
    lines = [
        '/* ==========================================================================',
        '   EO SELECTOR OPTIONS -- generated by 0_tools/buildEoOptions.py from the',
        '   consolidated EO source table. Do not edit by hand; re-run the script.',
        '   Loaded by openEoSection() (lcars.js) before eoPanel.js.',
        '   ========================================================================== */',
        '',
    ]
    for constant, values in options.items():
        lines.append(f'const {constant} = [')
        lines.extend(f'    {json.dumps(value)},' for value in values)
        lines.append('];')
        lines.append('')
    lines.append('// One entry per source in the table: the label shown in SOURCE and the lists the six filters match against.')
    lines.append('const EO_SOURCES = [')
    lines.extend(f'    {json.dumps(source, ensure_ascii=False)},' for source in sources)
    lines.append('];')
    lines.append('')
    return '\n'.join(lines)


if __name__ == '__main__':
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    built = buildOptions(xlsx)
    with open(OUT, 'w', encoding='utf-8', newline='\n') as f:
        f.write(render(built, buildSources(xlsx)))
    print(f'{OUT}: ' + ', '.join(f'{len(v)} {k}' for k, v in built.items()) + f', {len(buildSources(xlsx))} sources')
